import argparse
import asyncio
import io
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

URL = (
    "https://amis-rapid.uplb.edu.ph/"
    "rapid-tools/ocs/enrolled_student_summary_list.php"
    "?orderby=alast_name"
)

CDP_ENDPOINT = "http://127.0.0.1:9222"

OUTPUT_FOLDER = Path("downloaded_form5")
OUTPUT_FOLDER.mkdir(exist_ok=True)

EXCEL_PATH = OUTPUT_FOLDER / "form5_tracking.xlsx"
SHEET_NAME = "Form 5 Tracking"

# How long to wait for the form5.php response before giving up (seconds)
RESPONSE_TIMEOUT = 8.0

# Small pacing delay between students (seconds). Lower = faster,
# but too low can overload the AMIS server or trigger rate limiting.
INTER_STUDENT_DELAY = 0.15

# ------------------------------------------------------------
# Registration status = the watermark stamped across the PDF.
#
# The watermark is extracted as letter-spaced text, e.g.
# "O F F I C I A L L Y  R E G I S T E R E D", and appears before the
# "UP FORM 5." header. Keys below are written WITHOUT spaces because
# they are compared against whitespace-stripped ("squeezed") text.
# Checked top to bottom; first match wins.
# ------------------------------------------------------------
WATERMARK_TAGS = [
    ("OFFICIALLYREGISTERED", "Officially Registered"),
    ("BILLING", "Billing"),
]
DEFAULT_REGISTRATION_LABEL = "Unknown"

# ------------------------------------------------------------
# Scholarship / privilege = the value printed in the
# "SCHOLARSHIP / PRIVILEGES" box (bottom-right of the form),
# e.g. "RA 10931 FREE TUITION" or "PD80".
#
# Only that box is searched -- NOT the whole PDF, because every
# Form 5 contains the sentence "...to avail Free Tuition and Other
# School Fees", which would otherwise tag everyone as RA 10931.
#
# Keys are no-space, uppercase. Longer codes must come before
# shorter codes they contain (FDS before FD). If the box has a value
# that matches none of these, the raw value is written instead.
# ------------------------------------------------------------
SCHOLARSHIP_TAGS = [
    ("RA10931", "RA 10931 - Free Tuition"),
    ("FREETUITION", "RA 10931 - Free Tuition"),
    ("FDS", "FDS"),
    ("FD", "FD"),
    ("TFE", "TFE"),
    ("PD33", "PD 33"),
    ("PD60", "PD 60"),
    ("PD80", "PD 80"),
]
DEFAULT_SCHOLARSHIP_LABEL = "NE"


# ============================================================
# PDF PARSING
# ============================================================

NAME_RE = re.compile(r"\bNAME\s*:\s*(.+)", re.IGNORECASE)

# Form 5 prints e.g. "STUDENT NO. 202310846" (no dash). Also accept
# the dashed "2023-10846" style just in case.
STUDENT_NO_RE = re.compile(r"STUDENT\s*NO\.?\s*:?\s*(\d{4}-?\d{5})", re.IGNORECASE)

# "COLLEGE PROGRAM TERM & SY\nCEAT BSCE" -> program is the 2nd token
PROGRAM_RE = re.compile(r"COLLEGE\s+PROGRAM\s+TERM\s*&\s*SY\s+(\S+)\s+(\S+)", re.IGNORECASE)

SCHOLARSHIP_BLOCK_RE = re.compile(
    r"SCHOLARSHIP\s*/\s*PRIVILEGES\s*(.*?)\s*(?:Change of Matriculation|Deposit Fee|Date Generated|$)",
    re.IGNORECASE | re.DOTALL,
)

# Fallback for the degree when the PDF layout is unexpected: look in
# the student-list row text. Case-sensitive so surnames don't match.
ROW_DEGREE_RE = re.compile(r"\b(?:BS|BA|MS|MA|PHD)[A-Z]{1,10}(?:_[A-Z0-9]+)?\b")


@dataclass
class Form5Info:
    name: str | None
    student_number: str | None
    degree: str | None
    registration_status: str
    scholarship: str


def squeeze(text: str) -> str:
    """Strip ALL whitespace and uppercase, so letter-spaced stamps match."""
    return re.sub(r"\s+", "", text or "").upper()


def clean_filename(name: str) -> str:
    """Make the student name safe for a Windows filename."""
    name = re.sub(r'[<>:"/\\|?*]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name.strip()


def normalize_text(s: str) -> str:
    """Lowercase + collapse whitespace, for fuzzy matching."""
    return re.sub(r"\s+", " ", s or "").strip().lower()


def pdf_text(source) -> str:
    """Extract all text from a PDF (a path or raw bytes)."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    try:
        reader = PdfReader(source)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        print(f"ERROR extracting PDF text: {e}")
        return ""


def extract_student_name(text: str) -> str | None:
    match = NAME_RE.search(text)
    if not match:
        return None
    return clean_filename(match.group(1).split("\n")[0]) or None


def extract_student_number(text: str) -> str | None:
    match = STUDENT_NO_RE.search(text)
    return match.group(1) if match else None


def extract_degree(text: str, row_text: str = "") -> str | None:
    match = PROGRAM_RE.search(text)
    if match:
        return clean_filename(match.group(2))
    match = ROW_DEGREE_RE.search(row_text or "")
    return clean_filename(match.group(0)) if match else None


def extract_registration_status(text: str) -> str:
    """Read the watermark (OFFICIALLY REGISTERED / BILLING)."""
    # Normally the watermark is everything before the form header.
    header_pos = text.upper().find("UP FORM 5")
    if header_pos > 0:
        region = squeeze(text[:header_pos])
        for key, label in WATERMARK_TAGS:
            if key in region:
                return label

    # Fallback: find the letter-spaced stamp anywhere. Requiring
    # whitespace between every letter means ordinary words never match.
    for key, label in WATERMARK_TAGS:
        if re.search(r"\s+".join(key), text, re.IGNORECASE):
            return label

    return DEFAULT_REGISTRATION_LABEL


def extract_scholarship(text: str) -> str:
    """Read the value in the SCHOLARSHIP / PRIVILEGES box."""
    match = SCHOLARSHIP_BLOCK_RE.search(text)
    if not match:
        return DEFAULT_SCHOLARSHIP_LABEL

    raw = re.sub(r"\s+", " ", match.group(1)).strip()
    if not raw:
        return DEFAULT_SCHOLARSHIP_LABEL

    squeezed = squeeze(raw)
    for key, label in SCHOLARSHIP_TAGS:
        if key in squeezed:
            return label

    # Unknown code: keep the raw value (if it looks like a short code)
    # so it can be added to SCHOLARSHIP_TAGS later.
    return raw if len(raw) <= 40 else DEFAULT_SCHOLARSHIP_LABEL


def parse_form5(text: str, row_text: str = "") -> Form5Info:
    return Form5Info(
        name=extract_student_name(text),
        student_number=extract_student_number(text),
        degree=extract_degree(text, row_text),
        registration_status=extract_registration_status(text),
        scholarship=extract_scholarship(text),
    )


# ============================================================
# FILES / DEDUP
# ============================================================

def unique_filename(folder: Path, name: str) -> Path:
    """Prevent overwriting existing files."""
    path = folder / f"{name}.pdf"
    counter = 2
    while path.exists():
        path = folder / f"{name} ({counter}).pdf"
        counter += 1
    return path


def make_key(degree: str | None, student_name: str) -> str:
    """Canonical dedup key for a student's Form 5 (matches the filename)."""
    base = f"{degree}_{student_name}" if degree else student_name
    return normalize_text(base)


def build_existing_index(folder: Path) -> dict:
    """
    Scan already-downloaded PDFs so we can:
      (a) skip re-downloading a student we already have (post-check), and
      (b) recognize a student's row BEFORE clicking, by checking whether
          their name (pulled from an existing filename) shows up in the
          row text (pre-check).
    """
    keys = set()
    names = set()

    for f in folder.glob("*.pdf"):
        if f.name.startswith(("_temporary_", "UNKNOWN_")):
            continue

        stem = re.sub(r"\s*\(\d+\)$", "", f.stem)  # strip " (2)" style suffix
        keys.add(normalize_text(stem))

        parts = stem.split("_", 1)
        norm = normalize_text(parts[1] if len(parts) == 2 else stem)
        if norm:
            names.add(norm)

    return {"keys": keys, "names": names}


# ============================================================
# EXCEL TRACKING SHEET
# ============================================================

EXCEL_HEADERS = [
    "Date Processed",
    "Student Number",
    "Student Name",
    "Degree/Program",
    "Registration Status",
    "Scholarship Privilege",
    "PDF Filename",
    "Notes",
]
EXCEL_WIDTHS = [16, 16, 30, 18, 22, 24, 34, 30]


def init_excel(path: Path, fresh: bool = False):
    """Open the tracking workbook, creating it (or just the sheet) if needed.
    fresh=True replaces the tracking sheet with an empty one."""
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()

    if not path.exists():
        wb.active.title = SHEET_NAME
        _write_excel_header(wb.active)
        return wb, wb.active

    if fresh and SHEET_NAME in wb.sheetnames:
        position = wb.sheetnames.index(SHEET_NAME)
        wb.remove(wb[SHEET_NAME])
        ws = wb.create_sheet(SHEET_NAME, position)
        _write_excel_header(ws)
    elif SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_NAME)
        _write_excel_header(ws)
    else:
        ws = wb[SHEET_NAME]
    return wb, ws


def _write_excel_header(ws):
    header_fill = PatternFill(start_color="8D1436", end_color="8D1436", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")

    for col_idx, header in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"

    for col_idx, width in enumerate(EXCEL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def append_excel_row(ws, info: Form5Info, filename: str, notes: str = "", processed: datetime | None = None):
    row_idx = ws.max_row + 1
    values = [
        (processed or datetime.now()).strftime("%Y-%m-%d %H:%M"),
        info.student_number or "",
        info.name or "UNKNOWN",
        info.degree or "",
        info.registration_status,
        info.scholarship,
        filename,
        notes,
    ]
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = Font(name="Arial")
    # Keep student numbers as text so Excel doesn't reformat them
    ws.cell(row=row_idx, column=2).number_format = "@"


def excel_key(student_number: str | None, student_name: str | None) -> str:
    return normalize_text(student_number) if student_number else normalize_text(student_name)


def build_excel_existing_keys(ws) -> set:
    """Avoid duplicate Excel rows across runs (by student number, else name)."""
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row:
            key = excel_key(row[1], row[2])
            if key:
                keys.add(key)
    return keys


def save_excel(wb, path: Path):
    """Save, and if Excel has the file open/locked, ask the user to close it."""
    while True:
        try:
            wb.save(path)
            return
        except PermissionError:
            input(f"\nCannot save '{path}' - it is probably open in Excel. "
                  "Close it, then press ENTER to retry...")


def rebuild_excel():
    """Re-parse every PDF in OUTPUT_FOLDER and rewrite the tracking sheet.
    Notes typed into the old sheet are kept (matched by PDF filename)."""
    old_notes = {}
    if EXCEL_PATH.exists():
        wb = openpyxl.load_workbook(EXCEL_PATH)
        if SHEET_NAME in wb.sheetnames:
            for row in wb[SHEET_NAME].iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 8 and row[6] and row[7]:
                    old_notes[row[6]] = row[7]

    wb, ws = init_excel(EXCEL_PATH, fresh=True)
    pdfs = sorted(f for f in OUTPUT_FOLDER.glob("*.pdf") if not f.name.startswith("_temporary_"))

    for f in pdfs:
        info = parse_form5(pdf_text(f))
        notes = old_notes.get(f.name, "")
        if not info.name and not notes:
            notes = "Name could not be parsed from PDF - please check manually"
        append_excel_row(ws, info, f.name, notes, datetime.fromtimestamp(f.stat().st_mtime))
        print(f"{info.student_number or '?':>10}  {info.registration_status:<22}  "
              f"{info.scholarship:<24}  {f.name}")

    save_excel(wb, EXCEL_PATH)
    print(f"\nRebuilt '{EXCEL_PATH}' from {len(pdfs)} PDF(s).")


# ============================================================
# BROWSER HELPERS
# ============================================================

async def close_extra_tabs(context, main_page):
    """Close all tabs except the main AMIS page. Re-check a few times
    because AMIS may open PDF tabs asynchronously."""
    for _ in range(5):
        await asyncio.sleep(0.3)
        extra_pages = [p for p in context.pages if p != main_page]
        if not extra_pages:
            return
        for extra_page in extra_pages:
            try:
                await extra_page.close()
            except Exception as e:
                print(f"  Could not close tab: {e}")


async def click_and_capture_form5(page, context, button):
    """Click a Form 5 button and return the form5.php response (or None).
    Works whether AMIS loads the PDF in the same tab or a new one."""
    form5_response = None
    form5_event = asyncio.Event()

    def handle_response(response):
        nonlocal form5_response
        if "form5.php" in response.url and not form5_event.is_set():
            form5_response = response
            form5_event.set()

    def handle_new_page(new_page):
        new_page.on("response", handle_response)

    page.on("response", handle_response)
    context.on("page", handle_new_page)
    try:
        await button.click(timeout=10000)
        await asyncio.wait_for(form5_event.wait(), timeout=RESPONSE_TIMEOUT)
    except asyncio.TimeoutError:
        pass
    except Exception as e:
        print(f"Could not click Form 5: {e}")
    finally:
        page.remove_listener("response", handle_response)
        context.remove_listener("page", handle_new_page)

    return form5_response


# ============================================================
# MAIN
# ============================================================

async def download_all(debug: bool = False):
    from playwright.async_api import async_playwright

    async with async_playwright() as p:

        print("=" * 70)
        print("[+] FORM THUNDER // AMIS DOWNLOADER")
        print("=" * 70)

        # Connect to the already-open, already-logged-in Edge (see CLAUDE.md)
        browser = await p.chromium.connect_over_cdp(CDP_ENDPOINT)
        context = browser.contexts[0]
        page = context.pages[0]

        print("\nOpening AMIS...")
        await page.goto(URL, wait_until="domcontentloaded")
        print("\nIf AMIS requires login, log in now.")
        input("Press ENTER after the student list is visible...")

        existing_index = build_existing_index(OUTPUT_FOLDER)
        print(f"\nFound {len(existing_index['keys'])} existing Form 5 PDF(s) "
              f"in '{OUTPUT_FOLDER}'. These will be skipped.")

        excel_wb, excel_ws = init_excel(EXCEL_PATH)
        excel_existing_keys = build_excel_existing_keys(excel_ws)
        print(f"Tracking sheet: {EXCEL_PATH} ({len(excel_existing_keys)} row(s) already logged)")

        view_buttons = page.get_by_text("Form 5", exact=True)
        count = await view_buttons.count()
        print(f"\nFound {count} Form 5 buttons.")

        if count == 0:
            print("No Form 5 buttons were found. Inspect the page HTML and adjust the selector.")
            await browser.close()
            return

        skipped_early = skipped_duplicate = downloaded = failed = 0

        for index in range(count):
            print(f"\n--- STUDENT {index + 1} / {count} ---")

            view_buttons = page.get_by_text("Form 5", exact=True)
            if index >= await view_buttons.count():
                print("No more Form 5 buttons.")
                break

            button = view_buttons.nth(index)
            await button.scroll_into_view_if_needed()

            try:
                row = button.locator("xpath=ancestor::tr")
                row_text = await row.inner_text() if await row.count() else ""
            except Exception:
                row_text = ""

            # Pre-click duplicate check: skip without opening the PDF
            row_norm = normalize_text(row_text)
            matched = next((nm for nm in existing_index["names"] if nm in row_norm), None)
            if matched:
                print(f"SKIP (already downloaded): {matched}")
                skipped_early += 1
                continue

            form5_response = await click_and_capture_form5(page, context, button)

            pdf_bytes = b""
            if form5_response:
                try:
                    pdf_bytes = await form5_response.body()
                except Exception as e:
                    print(f"Could not read PDF: {e}")

            if not pdf_bytes.startswith(b"%PDF"):
                print("WARNING: form5.php was not detected or did not return a PDF.")
                failed += 1
                await close_extra_tabs(context, page)
                continue

            text = pdf_text(pdf_bytes)
            if debug:
                print("\nPDF TEXT:\n" + "-" * 60 + f"\n{text[:3000]}\n" + "-" * 60)

            info = parse_form5(text, row_text)
            print(f"Student No: {info.student_number or 'UNKNOWN'} | Degree: {info.degree or 'UNKNOWN'} | "
                  f"Status: {info.registration_status} | Scholarship: {info.scholarship}")

            if not info.name:
                unknown_file = unique_filename(OUTPUT_FOLDER, f"UNKNOWN_{index + 1}")
                unknown_file.write_bytes(pdf_bytes)
                print(f"WARNING: Could not determine student name. Saved as: {unknown_file}")
                failed += 1
                append_excel_row(excel_ws, info, unknown_file.name,
                                 "Name could not be parsed from PDF - please check manually")
                save_excel(excel_wb, EXCEL_PATH)

            elif make_key(info.degree, info.name) in existing_index["keys"]:
                print(f"SKIP (duplicate found after download): {info.name}")
                skipped_duplicate += 1

            else:
                filename = f"{info.degree}_{info.name}" if info.degree else info.name
                final_file = unique_filename(OUTPUT_FOLDER, filename)
                final_file.write_bytes(pdf_bytes)

                existing_index["keys"].add(make_key(info.degree, info.name))
                existing_index["names"].add(normalize_text(info.name))
                downloaded += 1
                print(f"SAVED: {final_file.name}")

                key = excel_key(info.student_number, info.name)
                if key not in excel_existing_keys:
                    append_excel_row(excel_ws, info, final_file.name)
                    save_excel(excel_wb, EXCEL_PATH)
                    excel_existing_keys.add(key)

            await close_extra_tabs(context, page)
            await page.wait_for_timeout(int(INTER_STUDENT_DELAY * 1000))

        save_excel(excel_wb, EXCEL_PATH)

        print("\n" + "=" * 70)
        print("FINISHED")
        print("=" * 70)
        print(f"Downloaded new:        {downloaded}")
        print(f"Skipped (pre-click):   {skipped_early}")
        print(f"Skipped (post-check):  {skipped_duplicate}")
        print(f"Failed:                {failed}")
        print(f"\nPDF files are in:\n{OUTPUT_FOLDER.absolute()}")
        print(f"Tracking sheet:\n{EXCEL_PATH.absolute()}")

        input("\nPress ENTER to close the browser...")
        await browser.close()


def main():
    parser = argparse.ArgumentParser(description="Download UPLB AMIS Form 5 PDFs and log them to Excel.")
    parser.add_argument("--rebuild-excel", action="store_true",
                        help="Re-parse the PDFs already in the output folder and rewrite the tracking sheet "
                             "(no browser needed).")
    parser.add_argument("--debug", action="store_true", help="Print the extracted PDF text for each student.")
    args = parser.parse_args()

    if args.rebuild_excel:
        rebuild_excel()
    else:
        asyncio.run(download_all(debug=args.debug))


if __name__ == "__main__":
    main()
