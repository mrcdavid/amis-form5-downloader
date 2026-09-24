"""Excel tracking sheet."""

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXCEL_HEADERS, EXCEL_WIDTHS, SHEET_NAME
from .models import Form5Info
from .text_utils import normalize_text


def init_excel(path: Path, fresh: bool = False):
    """Open the tracking workbook, creating it (or just the sheet) if needed.
    fresh=True replaces the tracking sheet with an empty one."""
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()

    if not path.exists():
        wb.active.title = SHEET_NAME
        _write_excel_header(wb.active)
        return wb, wb.active

    if fresh and SHEET_NAME in wb.sheetnames:
        position = wb.sheetnames.index(SHEET_NAME)
        wb.remove(wb[SHEET_NAME])
        ws = wb.create_sheet(SHEET_NAME, position)
        _write_excel_header(ws)
    elif SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_NAME)
        _write_excel_header(ws)
    else:
        ws = wb[SHEET_NAME]
    return wb, ws


def _write_excel_header(ws):
    header_fill = PatternFill(start_color="8D1436", end_color="8D1436", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")

    for col_idx, header in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"

    for col_idx, width in enumerate(EXCEL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def append_excel_row(ws, info: Form5Info, filename: str, notes: str = "", processed: datetime | None = None):
    row_idx = ws.max_row + 1
    values = [
        (processed or datetime.now()).strftime("%Y-%m-%d %H:%M"),
        info.student_number or "",
        info.name or "UNKNOWN",
        info.degree or "",
        info.registration_status,
        info.scholarship,
        filename,
        notes,
    ]
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = Font(name="Arial")
    # Keep student numbers as text so Excel doesn't reformat them
    ws.cell(row=row_idx, column=2).number_format = "@"


def excel_key(student_number: str | None, student_name: str | None) -> str:
    return normalize_text(student_number) if student_number else normalize_text(student_name)


def build_excel_existing_keys(ws) -> set:
    """Avoid duplicate Excel rows across runs (by student number, else name)."""
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row:
            key = excel_key(row[1], row[2])
            if key:
                keys.add(key)
    return keys


def save_excel(wb, path: Path):
    """Save, and if Excel has the file open/locked, ask the user to close it."""
    while True:
        try:
            wb.save(path)
            return
        except PermissionError:
            input(f"\nCannot save '{path}' - it is probably open in Excel. "
                  "Close it, then press ENTER to retry...")
