"""--rebuild-excel: re-parse downloaded PDFs and rewrite the sheet (no browser)."""

from datetime import datetime

import openpyxl

from .config import EXCEL_PATH, OUTPUT_FOLDER, SHEET_NAME
from .excel import append_excel_row, init_excel, save_excel
from .parsing import parse_form5, pdf_text


def rebuild_excel():
    """Re-parse every PDF in OUTPUT_FOLDER and rewrite the tracking sheet.
    Notes typed into the old sheet are kept (matched by PDF filename)."""
    old_notes = {}
    if EXCEL_PATH.exists():
        wb = openpyxl.load_workbook(EXCEL_PATH)
        if SHEET_NAME in wb.sheetnames:
            for row in wb[SHEET_NAME].iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 8 and row[6] and row[7]:
                    old_notes[row[6]] = row[7]

    wb, ws = init_excel(EXCEL_PATH, fresh=True)
    pdfs = sorted(f for f in OUTPUT_FOLDER.glob("*.pdf") if not f.name.startswith("_temporary_"))

    for f in pdfs:
        info = parse_form5(pdf_text(f))
        notes = old_notes.get(f.name, "")
        if not info.name and not notes:
            notes = "Name could not be parsed from PDF - please check manually"
        append_excel_row(ws, info, f.name, notes, datetime.fromtimestamp(f.stat().st_mtime))
        print(f"{info.student_number or '?':>10}  {info.registration_status:<22}  "
              f"{info.scholarship:<24}  {f.name}")

    save_excel(wb, EXCEL_PATH)
    print(f"\nRebuilt '{EXCEL_PATH}' from {len(pdfs)} PDF(s).")
